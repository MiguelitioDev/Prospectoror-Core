from pathlib import Path
import sqlite3
import openpyxl
import pytest

from prospector import cli, config, historico, ibge, osm, sites
from prospector.cli import main
from prospector.historico import municipios_varridos, pendentes
from prospector.modelos import Lead, Municipio


def ler_ids(caminho: Path) -> list[str]:
    """Lê todos os IDs presentes na aba Abordar da planilha informada."""
    wb = openpyxl.load_workbook(Path(caminho), read_only=True, data_only=True)
    try:
        ws = wb["Abordar"]
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            return []
        cab = [str(c or "").strip() for c in linhas[0]]
        if "ID" not in cab:
            return []
        idx = cab.index("ID")
        return [
            str(r[idx]).strip()
            for r in linhas[1:]
            if r and idx < len(r) and r[idx] is not None and str(r[idx]).strip()
        ]
    finally:
        wb.close()


class _AmbienteCLI:
    def __init__(self, tmp_path: Path):
        self.tmp_path = tmp_path
        self.dados = tmp_path / "dados"
        self.saida = tmp_path / "saida"
        self.saida.mkdir(parents=True, exist_ok=True)
        self.municipios = [
            Municipio(ibge="2304400", nome="Fortaleza", uf="CE", lat=-3.73, lng=-38.52, populacao=2574412),
            Municipio(ibge="2303709", nome="Caucaia", uf="CE", lat=-3.736, lng=-38.653, populacao=355679),
        ]
        self.n_municipios = len(self.municipios)
        self.chamadas_osm = 0
        self._falhas_osm: set[str] = set()
        self._leads_por_mun: dict[str, list[Lead]] = {
            "2304400": [
                Lead(
                    id="osm:node/1",
                    fonte="OSM",
                    nome="Pousada Iracema",
                    tipo="Pousada/Hotel",
                    municipio_ibge="2304400",
                    cidade="Fortaleza",
                    uf="CE",
                    lat=-3.73,
                    lng=-38.52,
                    telefone="+5585999991111",
                    whatsapp="https://wa.me/5585999991111",
                    status_site="Sem site",
                )
            ],
            "2303709": [
                Lead(
                    id="osm:node/2",
                    fonte="OSM",
                    nome="Pousada Cumbuco",
                    tipo="Pousada/Hotel",
                    municipio_ibge="2303709",
                    cidade="Caucaia",
                    uf="CE",
                    lat=-3.736,
                    lng=-38.653,
                    telefone="+5585999992222",
                    whatsapp="https://wa.me/5585999992222",
                    status_site="Sem site",
                )
            ],
        }
        self._conexoes: list = []

    def buscar_osm(self, mun: Municipio, tipos: dict | None = None) -> list[Lead] | None:
        self.chamadas_osm += 1
        if mun.ibge in self._falhas_osm:
            return None
        originais = self._leads_por_mun.get(mun.ibge, [])
        return [Lead.de_dict(l.para_dict()) for l in originais]

    def osm_devolve_os_mesmos_mais_um_novo(self) -> None:
        self._leads_por_mun["2304400"].append(
            Lead(
                id="osm:node/999",
                fonte="OSM",
                nome="Pousada Nova",
                tipo="Pousada/Hotel",
                municipio_ibge="2304400",
                cidade="Fortaleza",
                uf="CE",
                lat=-3.73,
                lng=-38.52,
                telefone="+5585999999999",
                whatsapp="https://wa.me/5585999999999",
                status_site="Sem site",
            )
        )

    def osm_falha_em(self, ibge_cod: str) -> None:
        self._falhas_osm.add(str(ibge_cod))

    def osm_devolve_n_leads_com_contato(self, n: int) -> None:
        self._leads_por_mun = {"2304400": [], "2303709": []}
        for i in range(1, n + 1):
            mun_alvo = "2304400" if i % 2 == 1 else "2303709"
            cidade = "Fortaleza" if mun_alvo == "2304400" else "Caucaia"
            self._leads_por_mun[mun_alvo].append(
                Lead(
                    id=f"osm:node/{i}",
                    fonte="OSM",
                    nome=f"Empresa {i}",
                    tipo="Pousada/Hotel",
                    municipio_ibge=mun_alvo,
                    cidade=cidade,
                    uf="CE",
                    lat=-3.73,
                    lng=-38.52,
                    telefone=f"+558599999{i:04d}",
                    whatsapp=f"https://wa.me/558599999{i:04d}",
                    status_site="Sem site",
                )
            )

    def osm_devolve(self, nome: str, uf: str = "CE", telefone: str = "+558532221111") -> None:
        self._leads_por_mun["2304400"].append(
            Lead(
                id="osm:node/777",
                fonte="OSM",
                nome=nome,
                tipo="Pousada/Hotel",
                municipio_ibge="2304400",
                cidade="Fortaleza",
                uf=uf,
                lat=-3.73,
                lng=-38.52,
                telefone=telefone,
                status_site="Sem site",
            )
        )

    def planilhas(self) -> list[Path]:
        return sorted(self.saida.glob("leads_*.xlsx"), key=lambda p: (p.stat().st_mtime_ns, p.name))

    def unica_planilha(self) -> Path:
        arqs = self.planilhas()
        assert len(arqs) == 1
        return arqs[0]

    def planilha_mais_recente(self) -> Path:
        arqs = self.planilhas()
        assert arqs
        return arqs[-1]

    def conn(self):
        c = historico.abrir(self.dados / "prospector.db")
        self._conexoes.append(c)
        return c

    def fechar(self) -> None:
        for c in self._conexoes:
            try:
                c.close()
            except Exception:
                pass

    def preencher_status(self, caminho: Path, linha: int, status: str) -> None:
        wb = openpyxl.load_workbook(caminho)
        ws = wb["Abordar"]
        cab = [c.value for c in ws[1]]
        col_status = cab.index("Status contato") + 1
        col_id = cab.index("ID") + 1
        self._ultimo_id_editado = str(ws.cell(row=linha + 1, column=col_id).value)
        ws.cell(row=linha + 1, column=col_status, value=status)
        wb.save(caminho)
        wb.close()

    def lead_bloqueado(self, linha: int = 1) -> bool:
        lead_id = getattr(self, "_ultimo_id_editado", None)
        if not lead_id:
            ids = ler_ids(self.planilha_mais_recente())
            lead_id = ids[linha - 1]
        c = self.conn()
        row = c.execute("SELECT bloqueado FROM leads WHERE id = ?", (lead_id,)).fetchone()
        return bool(row and row[0] == 1)

    def copiar_planilha_antiga_com(self, nome: str, cidade: str, telefone: str) -> Path:
        caminho = self.saida / "leads_2026-01-01_0000.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(["Nome", "Cidade", "Telefone", "E-mail", "WhatsApp", "Redes sociais"])
        ws.append([nome, cidade, telefone, "", "", ""])
        wb.save(caminho)
        wb.close()
        return caminho

    def nomes_exportados(self) -> list[str]:
        caminho = self.planilha_mais_recente()
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        try:
            ws = wb["Abordar"]
            linhas = list(ws.iter_rows(values_only=True))
            if not linhas:
                return []
            cab = [str(c or "").strip() for c in linhas[0]]
            idx = cab.index("Nome")
            return [str(r[idx]).strip() for r in linhas[1:] if r and idx < len(r) and r[idx]]
        finally:
            wb.close()


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    """Configura ambiente isolado em tmp_path com fakes de IBGE, OSM, verificação de sites e sleep."""
    amb = _AmbienteCLI(tmp_path)
    monkeypatch.setattr(config, "PASTA_DADOS", str(amb.dados))
    monkeypatch.setattr(config, "PASTA_SAIDA", str(amb.saida))
    monkeypatch.setattr(ibge, "listar_municipios", lambda ufs, cache_dir: list(amb.municipios))
    monkeypatch.setattr(osm, "buscar_municipio", amb.buscar_osm)
    monkeypatch.setattr(sites, "site_no_ar", lambda url: True)
    monkeypatch.setattr(cli.time, "sleep", lambda s: None)
    try:
        yield amb
    finally:
        amb.fechar()


def test_segunda_rodada_nao_repete(ambiente):
    assert main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)]) == 0
    primeira = ler_ids(ambiente.unica_planilha())
    assert primeira                                  # exportou algo
    ambiente.osm_devolve_os_mesmos_mais_um_novo()
    main(["buscar", "--uf", "CE", "--revarrer-dias", "0", "--saida", str(ambiente.saida)])
    segunda = ler_ids(ambiente.planilha_mais_recente())
    assert set(segunda) == {"osm:node/999"}


def test_falha_overpass_nao_marca_varrido(ambiente):
    ambiente.osm_falha_em("2304400")
    main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)])
    assert "2304400" not in municipios_varridos(ambiente.conn(), "osm", "2000-01-01T00:00:00")


def test_limite_deixa_resto_na_fila(ambiente):
    ambiente.osm_devolve_n_leads_com_contato(10)
    main(["buscar", "--uf", "CE", "--limite", "4", "--saida", str(ambiente.saida)])
    main(["buscar", "--uf", "CE", "--limite", "4", "--saida", str(ambiente.saida)])
    assert len(ler_ids(ambiente.planilha_mais_recente())) == 4
    assert ambiente.chamadas_osm == 1 * ambiente.n_municipios   # 2ª rodada não consultou de novo


def test_status_da_planilha_volta_e_bloqueia(ambiente):
    main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)])
    ambiente.preencher_status(ambiente.unica_planilha(), linha=1, status="Não contatar")
    main(["sincronizar", "--saida", str(ambiente.saida)])
    assert ambiente.lead_bloqueado(linha=1)


def test_planilha_antiga_sem_id_e_importada_como_legado(ambiente):
    ambiente.copiar_planilha_antiga_com(nome="Pousada Sol", cidade="Cumbuco - CE", telefone="(85) 3222-1111")
    ambiente.osm_devolve(nome="Pousada Sol", uf="CE", telefone="+558532221111")
    main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)])
    assert "Pousada Sol" not in ambiente.nomes_exportados()


def test_falha_ao_salvar_nao_marca_exportado(ambiente, monkeypatch):
    monkeypatch.setattr(cli, "gerar_planilha", lambda *a, **k: (_ for _ in ()).throw(OSError("disco cheio")))
    with pytest.raises(OSError):
        main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)])
    assert pendentes(ambiente.conn())               # continuam na fila


def test_planilha_ilegivel_nao_para_rodada(ambiente):
    (ambiente.saida / "leads_quebrada.xlsx").write_bytes(b"isto nao e um xlsx")
    assert main(["buscar", "--uf", "CE", "--saida", str(ambiente.saida)]) == 0


def test_buscar_cruza_com_cnpj_db_quando_existe(ambiente, fazer_estab):
    ambiente.dados.mkdir(parents=True, exist_ok=True)
    estab = fazer_estab(
        cnpj="99887766000155",
        nome="CAFE EXCLUSIVO RECEITA",
        tipo="Bar/Café",
        municipio_ibge="2304400",
        uf="CE",
        telefone="+5585988887777",
    )
    conn_cnpj = sqlite3.connect(ambiente.dados / "cnpj.db")
    try:
        conn_cnpj.execute(
            """
            CREATE TABLE estab (
                cnpj TEXT PRIMARY KEY, nome TEXT, cnae TEXT, tipo TEXT,
                municipio_ibge TEXT, uf TEXT, logradouro TEXT, bairro TEXT,
                cep TEXT, telefone TEXT, telefone2 TEXT, email TEXT, abertura TEXT
            )
            """
        )
        conn_cnpj.execute(
            "INSERT INTO estab VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                estab.cnpj, estab.nome, estab.cnae, estab.tipo, estab.municipio_ibge,
                estab.uf, estab.logradouro, estab.bairro, estab.cep, estab.telefone,
                estab.telefone2, estab.email, estab.abertura,
            ),
        )
        conn_cnpj.commit()
    finally:
        conn_cnpj.close()

    assert main(["buscar", "--uf", "CE", "--dados", str(ambiente.dados), "--saida", str(ambiente.saida)]) == 0
    wb = openpyxl.load_workbook(ambiente.unica_planilha(), read_only=True, data_only=True)
    try:
        ws = wb["Abordar"]
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip() for c in linhas[0]]
        idx_fonte = cab.index("Fonte")
        idx_id = cab.index("ID")
        fontes = {str(r[idx_id]): str(r[idx_fonte]) for r in linhas[1:] if r}
    finally:
        wb.close()

    assert fontes.get("cnpj:99887766000155") == "Receita"

