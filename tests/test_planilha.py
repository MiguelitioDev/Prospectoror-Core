import openpyxl
import pytest

from prospector.planilha import gerar_planilha, ler_legado, ler_status
from prospector.pontuacao import definir_entregas, pontuar


@pytest.fixture
def linhas_exemplo(fazer_lead):
    """Gera 3 linhas pontuadas e com entregas definidas para testes da planilha."""
    leads = [
        fazer_lead(
            id="osm:node/1",
            nome="Pousada Mar Azul",
            cidade="Fortaleza",
            uf="CE",
            telefone="+5585999991111",
            whatsapp="https://wa.me/5585999991111",
        ),
        fazer_lead(
            id="osm:node/2",
            nome="Restaurante Sabor",
            tipo="Restaurante",
            cidade="Caucaia",
            uf="CE",
            telefone="+5585999992222",
            whatsapp="https://wa.me/5585999992222",
        ),
        fazer_lead(
            id="osm:node/3",
            nome="Café da Praça",
            tipo="Bar/Café",
            cidade="Fortaleza",
            uf="CE",
            telefone="+558532223333",
        ),
    ]
    linhas = [pontuar(lead) for lead in leads]
    definir_entregas(linhas)
    return linhas


def test_planilha_ida_e_volta(tmp_path, linhas_exemplo):
    caminho = gerar_planilha(linhas_exemplo, [], {"Data": "2026-10-24"}, tmp_path / "leads.xlsx")
    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Abordar", "Resumo", "Rodada"]
    ws = wb["Abordar"]
    assert "Leads" in ws.tables
    assert ws.data_validations.dataValidation          # lista suspensa existe
    # usuário preenche, apaga a coluna Observações e uma linha fica sem ID
    cab = [c.value for c in ws[1]]
    ws.cell(row=2, column=cab.index("Status contato") + 1, value="Fechado")
    ws.cell(row=3, column=cab.index("ID") + 1).value = None
    ws.delete_cols(cab.index("Observações") + 1)
    wb.save(caminho)
    lidas = ler_status(caminho)
    assert lidas[0]["Status contato"] == "Fechado" and "Observações" not in lidas[0]
    assert len(lidas) == len(linhas_exemplo) - 1


def test_permission_error_salva_com_sufixo(tmp_path, linhas_exemplo, monkeypatch):
    original = openpyxl.Workbook.save

    def save(self, f):
        if str(f).endswith("leads.xlsx"):
            raise PermissionError
        original(self, f)

    monkeypatch.setattr(openpyxl.Workbook, "save", save)
    assert gerar_planilha(linhas_exemplo, [], {}, tmp_path / "leads.xlsx").name == "leads_2.xlsx"


def test_formato_antigo(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Leads"
    wb.active.append(["Nome", "Cidade"])
    wb.save(tmp_path / "v.xlsx")
    with pytest.raises(ValueError, match="formato antigo"):
        ler_status(tmp_path / "v.xlsx")


def test_abordar_vazio_e_pesquisar_com_linha(tmp_path, fazer_lead):
    sem_contato = pontuar(fazer_lead(id="osm:node/99", nome="Quiosque Sol", telefone="", whatsapp=""))
    caminho = gerar_planilha([], [sem_contato], {"Data": "2026-10-24"}, tmp_path / "so_pesquisar.xlsx")
    wb = openpyxl.load_workbook(caminho)
    try:
        assert "Pesquisar contato" in wb.sheetnames
        assert "Leads" not in wb["Abordar"].tables
    finally:
        wb.close()
    lidas = ler_status(caminho)
    assert len(lidas) == 1 and lidas[0]["ID"] == "osm:node/99"


def test_resumo_tem_bloco_por_cidade(tmp_path, linhas_exemplo):
    caminho = gerar_planilha(linhas_exemplo, [], {"Data": "2026-10-24"}, tmp_path / "resumo.xlsx")
    wb = openpyxl.load_workbook(caminho, data_only=True)
    try:
        ws = wb["Resumo"]
        linhas = [[cel.value for cel in row] for row in ws.iter_rows()]
    finally:
        wb.close()
    assert ["Cidade", "Quentes", "Mornos", "Empresas"] in linhas
    cidades_col = [r[0] for r in linhas if r and r[0]]
    assert "Fortaleza - CE" in cidades_col
