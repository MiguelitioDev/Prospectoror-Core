from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from prospector.config import STATUS_CONTATO

COLUNAS_ABORDAR = [
    "Chance de venda",
    "Prioridade",
    "Nome",
    "Tipo",
    "Cidade",
    "UF",
    "Preparar",
    "Protótipo",
    "Site final",
    "Horas estimadas",
    "Oferta inicial (R$)",
    "Valor alvo (R$)",
    "Mínimo (R$)",
    "Por quê",
    "Site",
    "URL atual",
    "Redes sociais",
    "Telefone",
    "WhatsApp",
    "Abrir conversa",
    "E-mail",
    "Distância (km)",
    "Abordagem",
    "Ver no Google Maps",
    "Nota Google (conferir)",
    "Status contato",
    "Valor fechado (R$)",
    "Observações",
    "Fonte",
    "CNPJ",
    "Aberta em",
    "ID",
]

COLUNAS_PESQUISAR = [
    "Nome",
    "Tipo",
    "Cidade",
    "UF",
    "Chance de venda",
    "Ver no Google Maps",
    "Telefone",
    "Observações",
    "ID",
]

_LARGURAS = {
    "Nome": 34,
    "Protótipo": 55,
    "Site final": 60,
    "Por quê": 60,
    "URL atual": 30,
    "Redes sociais": 30,
    "Ver no Google Maps": 30,
    "WhatsApp": 26,
    "Abrir conversa": 26,
    "Observações": 30,
}

_CORES_PRIORIDADE = {"Quente": "FECACA", "Morno": "FEF3C7", "Frio": "E5E7EB"}
_CORES_PREPARAR = {"Site final": "BBF7D0", "Protótipo": "BFDBFE", "Nada agora": "F3F4F6"}
_COLUNAS_LINK = ("Abrir conversa", "WhatsApp", "Ver no Google Maps", "URL atual")
_COLUNAS_WRAP = ("Protótipo", "Site final", "Por quê")
_COLUNAS_MOEDA = ("Oferta inicial (R$)", "Valor alvo (R$)", "Mínimo (R$)", "Valor fechado (R$)")
_FORMATO_MOEDA = '"R$" #,##0'


def _valor_celula(valor: Any) -> Any:
    """Converte estruturas não suportadas pelo openpyxl (dict/list/set/tuple) em string."""
    if valor is None:
        return ""
    if isinstance(valor, (dict, list, tuple, set)):
        return str(valor)
    return valor


def _limpar_linha(linha: dict) -> dict:
    """Remove chaves internas iniciadas com '_' e sanitiza valores para gravação."""
    if not isinstance(linha, dict):
        return {}
    return {str(k): _valor_celula(v) for k, v in linha.items() if not str(k).startswith("_")}


def _colunas_unicas(linhas: list[dict], padrao: list[str]) -> list[str]:
    """Extrai cabeçalhos únicos como strings a partir da primeira linha ou da lista padrão."""
    if not linhas:
        return list(padrao)
    vistos: set[str] = set()
    colunas: list[str] = []
    for k in linhas[0].keys():
        nome = str(k).strip()
        if not nome or nome.startswith("_") or nome in vistos:
            continue
        vistos.add(nome)
        colunas.append(nome)
    return colunas or list(padrao)


def _estilizar_cabecalho(ws) -> None:
    """Aplica fonte branca em negrito e fundo escuro na primeira linha da aba."""
    fonte = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F2937")
    for cel in ws[1]:
        cel.font = fonte
        cel.fill = fundo


def _aplicar_links(ws, colunas: list[str], nomes_colunas: tuple[str, ...]) -> None:
    """Transforma URLs não vazias nas colunas indicadas em hiperlinks clicáveis."""
    for nome_col in nomes_colunas:
        if nome_col not in colunas:
            continue
        idx = colunas.index(nome_col) + 1
        for row in range(2, ws.max_row + 1):
            cel = ws.cell(row=row, column=idx)
            url = str(cel.value or "").strip()
            if url:
                cel.hyperlink = url
                cel.style = "Hyperlink"


def _montar_aba_abordar(ws, abordar_limpo: list[dict]) -> None:
    """Popula e formata a aba principal 'Abordar' com tabela, cores, links e validação."""
    ws.title = "Abordar"
    colunas = _colunas_unicas(abordar_limpo, COLUNAS_ABORDAR)
    ws.append(colunas)
    for lead in abordar_limpo:
        ws.append([_valor_celula(lead.get(c, "")) for c in colunas])

    _estilizar_cabecalho(ws)

    if abordar_limpo:
        ultima_letra = get_column_letter(len(colunas))
        ultima_linha = len(abordar_limpo) + 1
        tabela = Table(displayName="Leads", ref=f"A1:{ultima_letra}{ultima_linha}")
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabela)

    if "Prioridade" in colunas:
        col_prio = colunas.index("Prioridade")
        for linha in ws.iter_rows(min_row=2, max_row=len(abordar_limpo) + 1):
            cel = linha[col_prio]
            cor = _CORES_PRIORIDADE.get(str(cel.value or ""))
            if cor:
                cel.fill = PatternFill("solid", fgColor=cor)

    if "Preparar" in colunas:
        col_prep = colunas.index("Preparar")
        for linha in ws.iter_rows(min_row=2, max_row=len(abordar_limpo) + 1):
            cel = linha[col_prep]
            cor = _CORES_PREPARAR.get(str(cel.value or ""))
            if cor:
                cel.fill = PatternFill("solid", fgColor=cor)
            cel.font = Font(bold=True)

    for nome_col in _COLUNAS_WRAP:
        if nome_col in colunas:
            letra = get_column_letter(colunas.index(nome_col) + 1)
            for cel in ws[letra][1:]:
                cel.alignment = Alignment(wrap_text=True, vertical="top")

    for nome_col in _COLUNAS_MOEDA:
        if nome_col in colunas:
            letra = get_column_letter(colunas.index(nome_col) + 1)
            for cel in ws[letra][1:]:
                cel.number_format = _FORMATO_MOEDA

    _aplicar_links(ws, colunas, _COLUNAS_LINK)

    if "Status contato" in colunas:
        letra_status = get_column_letter(colunas.index("Status contato") + 1)
        fim_dv = max(len(abordar_limpo) + 1, 2) + 500
        opcoes = ",".join(STATUS_CONTATO)
        dv = DataValidation(type="list", formula1=f'"{opcoes}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letra_status}2:{letra_status}{fim_dv}")

    for i, c in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _LARGURAS.get(c, 15)
    ws.freeze_panes = "D2"


def _montar_aba_pesquisar(wb: Workbook, pesquisar_limpo: list[dict]) -> None:
    """Cria a aba opcional 'Pesquisar contato' quando há leads sem canal de contato."""
    if not pesquisar_limpo:
        return
    ws = wb.create_sheet("Pesquisar contato")
    colunas = list(COLUNAS_PESQUISAR)
    ws.append(colunas)
    for lead in pesquisar_limpo:
        ws.append([_valor_celula(lead.get(c, "")) for c in colunas])

    _estilizar_cabecalho(ws)
    _aplicar_links(ws, colunas, ("Ver no Google Maps",))

    for i, c in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _LARGURAS.get(c, 15)


def _montar_aba_resumo(wb: Workbook, lista: list[dict]) -> None:
    """Cria a aba 'Resumo' com 4 blocos montados dinamicamente sem índices de linha fixos."""
    resumo = wb.create_sheet("Resumo")
    linhas_titulo: list[int] = []
    celulas_moeda: list[tuple[int, int]] = []

    def adicionar_titulo(colunas: list[str]) -> None:
        resumo.append(colunas)
        linhas_titulo.append(resumo.max_row)

    # Bloco 1: Prioridade
    adicionar_titulo(["Prioridade", "Empresas", "Soma do valor alvo (R$)", "Valor alvo médio (R$)"])
    for p in ("Quente", "Morno", "Frio"):
        grupo = [l for l in lista if l.get("Prioridade") == p]
        soma = sum(int(l.get("Valor alvo (R$)") or 0) for l in grupo)
        media = round(soma / len(grupo)) if grupo else 0
        resumo.append([p, len(grupo), soma, media])
        celulas_moeda.append((resumo.max_row, 3))
        celulas_moeda.append((resumo.max_row, 4))

    resumo.append([])

    # Bloco 2: Abordagem
    adicionar_titulo(["Abordagem", "Quentes", "Mornos"])
    for f in ("Presencial", "Ligação", "Mensagem"):
        quentes = sum(1 for l in lista if l.get("Abordagem") == f and l.get("Prioridade") == "Quente")
        mornos = sum(1 for l in lista if l.get("Abordagem") == f and l.get("Prioridade") == "Morno")
        resumo.append([f, quentes, mornos])

    resumo.append([])

    # Bloco 3: Preparar
    adicionar_titulo(["Preparar", "Empresas", "Horas estimadas", "Soma do valor alvo (R$)"])
    for p in ("Site final", "Protótipo", "Nada agora"):
        grupo = [l for l in lista if l.get("Preparar") == p]
        horas = sum(int(l.get("Horas estimadas") or 0) for l in grupo)
        soma = sum(int(l.get("Valor alvo (R$)") or 0) for l in grupo)
        resumo.append([p, len(grupo), horas, soma])
        celulas_moeda.append((resumo.max_row, 4))

    resumo.append([])

    # Bloco 4: Por cidade (top 15 por Quentes, depois Empresas)
    adicionar_titulo(["Cidade", "Quentes", "Mornos", "Empresas"])
    por_cidade: dict[str, list[int]] = {}
    for l in lista:
        cidade = str(l.get("Cidade") or "").strip()
        uf = str(l.get("UF") or "").strip()
        if not cidade and not uf:
            continue
        chave = f"{cidade} - {uf}" if (cidade and uf) else (cidade or uf)
        if chave not in por_cidade:
            por_cidade[chave] = [0, 0, 0]
        if l.get("Prioridade") == "Quente":
            por_cidade[chave][0] += 1
        elif l.get("Prioridade") == "Morno":
            por_cidade[chave][1] += 1
        por_cidade[chave][2] += 1

    ranking = sorted(
        por_cidade.items(),
        key=lambda item: (-item[1][0], -item[1][2], item[0]),
    )[:15]
    for nome_cidade, (quentes, mornos, total) in ranking:
        resumo.append([nome_cidade, quentes, mornos, total])

    for num_linha in linhas_titulo:
        for cel in resumo[num_linha]:
            cel.font = Font(bold=True)

    for r_idx, c_idx in celulas_moeda:
        resumo.cell(row=r_idx, column=c_idx).number_format = _FORMATO_MOEDA

    for col in "ABCD":
        resumo.column_dimensions[col].width = 24


def _montar_aba_rodada(wb: Workbook, info: dict) -> None:
    """Cria a aba 'Rodada' com os pares chave/valor dos metadados da execução."""
    ws = wb.create_sheet("Rodada")
    ws.append(["Campo", "Valor"])
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for k, v in (info or {}).items():
        ws.append([str(k), _valor_celula(v)])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40


def _salvar_com_fallback(wb: Workbook, caminho: Path) -> Path:
    """Salva o workbook criando a pasta-mãe; se houver PermissionError, tenta sufixos _2.._9."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(caminho)
        return caminho
    except PermissionError:
        for i in range(2, 10):
            alternativo = caminho.with_name(f"{caminho.stem}_{i}{caminho.suffix}")
            try:
                wb.save(alternativo)
                return alternativo
            except PermissionError:
                continue
        raise


def gerar_planilha(abordar: list[dict], pesquisar: list[dict], info: dict, caminho: Path) -> Path:
    """Gera a planilha de prospecção com abas Abordar, Pesquisar contato (se houver), Resumo e Rodada."""
    abordar_limpo = [_limpar_linha(l) for l in (abordar or [])]
    pesquisar_limpo = [_limpar_linha(l) for l in (pesquisar or [])]

    wb = Workbook()
    _montar_aba_abordar(wb.active, abordar_limpo)
    _montar_aba_pesquisar(wb, pesquisar_limpo)
    _montar_aba_resumo(wb, abordar_limpo)
    _montar_aba_rodada(wb, info or {})

    return _salvar_com_fallback(wb, Path(caminho))


def ler_status(caminho: Path) -> list[dict]:
    """Lê as abas 'Abordar' e 'Pesquisar contato' devolvendo ID e colunas de retorno preenchíveis."""
    wb = openpyxl.load_workbook(Path(caminho), read_only=True, data_only=True)
    try:
        nomes_abas = wb.sheetnames
        abas_alvo = [a for a in ("Abordar", "Pesquisar contato") if a in nomes_abas]
        colunas_retorno = ("Status contato", "Valor fechado (R$)", "Observações")
        resultado: list[dict] = []
        achou_coluna_id = False

        for nome_aba in abas_alvo:
            ws = wb[nome_aba]
            it = ws.iter_rows(values_only=True)
            cab_bruto = next(it, None)
            if not cab_bruto:
                continue
            cab = [str(c).strip() if c is not None else "" for c in cab_bruto]
            if "ID" not in cab:
                continue
            achou_coluna_id = True
            idx_id = cab.index("ID")
            indices_retorno = {col: cab.index(col) for col in colunas_retorno if col in cab}

            for row in it:
                if not row or all(v is None or str(v).strip() == "" for v in row):
                    continue
                val_id = row[idx_id] if idx_id < len(row) else None
                if val_id is None or str(val_id).strip() == "":
                    continue
                item: dict[str, Any] = {"ID": str(val_id).strip()}
                for col, idx in indices_retorno.items():
                    val = row[idx] if idx < len(row) else None
                    item[col] = "" if val is None else val
                resultado.append(item)

        if not achou_coluna_id and "Leads" in nomes_abas:
            raise ValueError("formato antigo")

        return resultado
    finally:
        wb.close()


def ler_legado(caminho: Path) -> list[dict]:
    """Lê a aba 'Leads' de uma planilha no formato antigo como lista de dicionários."""
    wb = openpyxl.load_workbook(Path(caminho), read_only=True, data_only=True)
    try:
        if "Leads" not in wb.sheetnames:
            return []
        ws = wb["Leads"]
        it = ws.iter_rows(values_only=True)
        cab_bruto = next(it, None)
        if not cab_bruto:
            return []
        cab = [str(c).strip() if c is not None else "" for c in cab_bruto]

        linhas: list[dict] = []
        for row in it:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            d: dict[str, Any] = {}
            for i, nome_col in enumerate(cab):
                if not nome_col:
                    continue
                val = row[i] if i < len(row) else None
                d[nome_col] = "" if val is None else val
            linhas.append(d)
        return linhas
    finally:
        wb.close()
